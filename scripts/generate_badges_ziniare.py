import os
import sys
import shutil
import pypdfium2
import docx
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas
from reportlab.graphics.shapes import Drawing, Rect, Circle, Line, Group, Polygon
from reportlab.graphics.barcode import qr

PHOTOS_DIR = os.path.join(os.getcwd(), "docs", "photos_traitees")

# 8 COMMERCIAUX - ZONE ZINIARÉ (AVEC PHOTOS TRAITÉES LORSQUE DISPONIBLES)
COMMERCIAUX_ZINIARE = [
    {
        'name': 'COMPAORE Nadiatou Lethycia',
        'slug': 'COMPAORE_Nadiatou_Lethycia',
        'code': 'COMPAORE7',
        'phone': '+226 74 28 37 60',
        'phone_display': '74 28 37 60',
        'zone': 'Ziniaré',
        'photo_path': os.path.join(PHOTOS_DIR, "photo_COMPAORE_Nadiatou_Lethycia.png")
    },
    {
        'name': 'SAVADOGO Rimnomma Cyprien',
        'slug': 'SAVADOGO_Rimnomma_Cyprien',
        'code': 'CYPRIEN226',
        'phone': '+226 64 87 76 46',
        'phone_display': '64 87 76 46',
        'zone': 'Ziniaré',
        'photo_path': os.path.join(PHOTOS_DIR, "photo_SAVADOGO_Rimnomma_Cyprien.png")
    },
    {
        'name': 'NACOULMA Judicaëla',
        'slug': 'NACOULMA_Judicaela',
        'code': 'NACOULMA226',
        'phone': '+226 06 94 00 92',
        'phone_display': '06 94 00 92',
        'zone': 'Ziniaré',
        'photo_path': os.path.join(PHOTOS_DIR, "photo_NACOULMA_Judicaela.png")
    },
    {
        'name': 'TENKODOGO Aminata',
        'slug': 'TENKODOGO_Aminata',
        'code': 'TENKODOG226',
        'phone': '+226 56 49 16 54',
        'phone_display': '56 49 16 54',
        'zone': 'Ziniaré',
        'photo_path': os.path.join(PHOTOS_DIR, "photo_TENKODOGO_Aminata.png")
    },
    {
        'name': 'YANOGO W. Christiane',
        'slug': 'YANOGO_W_Christiane',
        'code': 'YANOGO226',
        'phone': '+226 57 34 20 92',
        'phone_display': '57 34 20 92',
        'zone': 'Ziniaré',
        'photo_path': os.path.join(PHOTOS_DIR, "photo_YANOGO_W_Christiane.png")
    },
    {
        'name': 'OUEDRAOGO Stécie',
        'slug': 'OUEDRAOGO_Stecie',
        'code': 'STECIE226',
        'phone': '+226 61 30 60 03',
        'phone_display': '61 30 60 03',
        'zone': 'Ziniaré',
        'photo_path': os.path.join(PHOTOS_DIR, "photo_OUEDRAOGO_Stecie.png")
    },
    {
        'name': 'NACOULMA Issouf',
        'slug': 'NACOULMA_Issouf',
        'code': 'NACOULMA7',
        'phone': '+226 54 72 86 42',
        'phone_display': '54 72 86 42',
        'zone': 'Ziniaré',
        'photo_path': None
    },
    {
        'name': 'SANKARA Ella',
        'slug': 'SANKARA_Ella',
        'code': 'SANKARA226',
        'phone': '+226 51 40 81 10',
        'phone_display': '51 40 81 10',
        'zone': 'Ziniaré',
        'photo_path': None
    }
]

# Cas spécial : Si l'utilisateur souhaite ajouter le 9ème commercial (ILY Pouékomba Hyppolite)
COMMERCIAL_HYPPOLITE = {
    'name': 'ILY Pouékomba Hyppolite',
    'slug': 'ILY_Pouekomba_Hyppolite',
    'code': 'ILY226',
    'phone': '+226 77 92 30 85',
    'phone_display': '77 92 30 85',
    'zone': 'Ziniaré',
    'photo_path': os.path.join(PHOTOS_DIR, "photo_ILY_Pouekomba_Hyppolite.png")
}

def draw_logo(d, x, y, size=32):
    """Dessine le logo officiel FasoCarnet"""
    scale = size / 40.0
    g = Group()
    g.translate(x, y)
    g.scale(scale, scale)
    g.add(Rect(0, 0, 40, 40, rx=9, ry=9, fillColor=colors.HexColor('#064e3b'), strokeColor=None))
    g.add(Rect(5, 6, 5.5, 28, rx=2, ry=2, fillColor=colors.HexColor('#047857'), strokeColor=None))
    g.add(Rect(12, 6, 23, 28, rx=3, ry=3, fillColor=colors.white, strokeColor=None))
    g.add(Line(16, 25, 31, 25, strokeColor=colors.HexColor('#e2e8f0'), strokeWidth=2))
    g.add(Line(16, 19, 26, 19, strokeColor=colors.HexColor('#e2e8f0'), strokeWidth=2))
    g.add(Circle(25, 13, 5, fillColor=colors.HexColor('#d97706'), strokeColor=None))
    star_pts = [29, 30, 31, 28, 33, 28, 31.5, 26.5, 32.5, 24.5, 30.5, 26, 28.5, 24.5, 29.5, 26.5, 28, 28, 30, 28]
    g.add(Polygon(star_pts, fillColor=colors.HexColor('#ef4444'), strokeColor=None))
    d.add(g)

def draw_qr_code(c, url, x, y, size_pt):
    widget = qr.QrCodeWidget(url)
    bounds = widget.getBounds()
    w = bounds[2] - bounds[0]
    h = bounds[3] - bounds[1]
    d = Drawing(size_pt, size_pt, transform=[size_pt / w, 0, 0, size_pt / h, 0, 0])
    d.add(widget)
    d.drawOn(c, x, y)

def draw_badge_recto(c, x, y, w, h, commercial):
    c.saveState()
    # 1. Fond blanc & bordure
    c.setFillColor(colors.white)
    c.setStrokeColor(colors.HexColor('#064e3b'))
    c.setLineWidth(1.2)
    c.roundRect(x, y, w, h, 4.5 * mm, fill=1, stroke=1)
    
    # 2. Bandeau supérieur officiel (Émeraude profonde)
    header_h = 17.5 * mm
    c.setFillColor(colors.HexColor('#064e3b'))
    c.roundRect(x + 0.6, y + h - header_h - 0.6, w - 1.2, header_h, 4 * mm, fill=1, stroke=0)
    c.rect(x + 0.6, y + h - header_h - 0.6, w - 1.2, 6 * mm, fill=1, stroke=0)
    
    # Accent Or
    c.setFillColor(colors.HexColor('#d97706'))
    c.rect(x + 0.6, y + h - header_h - 1.4 * mm, w - 1.2, 1.2 * mm, fill=1, stroke=0)
    
    # Logo & Titre
    logo_size = 20
    logo_d = Drawing(logo_size, logo_size)
    draw_logo(logo_d, 0, 0, size=logo_size)
    logo_d.drawOn(c, x + 3.2 * mm, y + h - 14.5 * mm)
    
    c.setFont("Helvetica-Bold", 10.5)
    c.setFillColor(colors.white)
    c.drawString(x + 12.8 * mm, y + h - 8.5 * mm, "FasoCarnet")
    
    c.setFont("Helvetica-Bold", 5.2)
    c.setFillColor(colors.HexColor('#fef08a'))
    c.drawString(x + 12.8 * mm, y + h - 11.4 * mm, "CAISSE & CARNET DIGITAL")
    
    c.setFont("Helvetica", 4.8)
    c.setFillColor(colors.HexColor('#a7f3d0'))
    c.drawString(x + 12.8 * mm, y + h - 14 * mm, "RESEAU COMMERCIAL OFFICIEL")
    
    # 3. Mention de la carte
    c.setFont("Helvetica-Bold", 6.2)
    c.setFillColor(colors.HexColor('#064e3b'))
    c.drawCentredString(x + w / 2, y + 64 * mm, "CARTE D'ACCREDITATION TERRAIN")
    
    # 4. Cadre photo (Photo traitée réelle ou Silhouette stylisée)
    photo_w = 21 * mm
    photo_h = 24 * mm
    photo_x = x + (w - photo_w) / 2
    photo_y = y + 37.5 * mm
    
    photo_path = commercial.get('photo_path')
    if photo_path and os.path.exists(photo_path):
        # Affichage photo réelle avec coins arrondis
        c.saveState()
        clip_p = c.beginPath()
        clip_p.roundRect(photo_x, photo_y, photo_w, photo_h, 2.5 * mm)
        c.clipPath(clip_p, stroke=0, fill=0)
        c.drawImage(photo_path, photo_x, photo_y, photo_w, photo_h, preserveAspectRatio=False)
        c.restoreState()
        # Bordure soignée
        c.setStrokeColor(colors.HexColor('#047857'))
        c.setLineWidth(1.0)
        c.roundRect(photo_x, photo_y, photo_w, photo_h, 2.5 * mm, fill=0, stroke=1)
    else:
        # Silhouette par défaut
        c.setFillColor(colors.HexColor('#f8fafc'))
        c.setStrokeColor(colors.HexColor('#cbd5e1'))
        c.setLineWidth(0.8)
        c.roundRect(photo_x, photo_y, photo_w, photo_h, 2.5 * mm, fill=1, stroke=1)
        
        c.setFillColor(colors.HexColor('#94a3b8'))
        c.circle(photo_x + photo_w / 2, photo_y + photo_h - 8 * mm, 4 * mm, fill=1, stroke=0)
        c.roundRect(photo_x + 2.5 * mm, photo_y + 1.5 * mm, photo_w - 5 * mm, 9 * mm, 3 * mm, fill=1, stroke=0)
        
        c.setFont("Helvetica-Bold", 4.5)
        c.setFillColor(colors.HexColor('#64748b'))
        c.drawCentredString(photo_x + photo_w / 2, photo_y + 2.4 * mm, "PHOTO AGENT")
    
    # 5. Nom & Titre
    name_str = commercial['name'].upper()
    font_size_name = 8.5 if len(name_str) > 22 else 9.5
    c.setFont("Helvetica-Bold", font_size_name)
    c.setFillColor(colors.HexColor('#0f172a'))
    c.drawCentredString(x + w / 2, y + 33 * mm, name_str)
    
    c.setFont("Helvetica-Bold", 6)
    c.setFillColor(colors.HexColor('#047857'))
    c.drawCentredString(x + w / 2, y + 29.5 * mm, "CONSEILLER COMMERCIAL TERRAIN")
    
    # 6. Bloc Infos
    info_box_y = y + 8.8 * mm
    info_box_h = 18.5 * mm
    info_box_w = w - 6 * mm
    info_box_x = x + 3 * mm
    
    c.setFillColor(colors.HexColor('#f8fafc'))
    c.setStrokeColor(colors.HexColor('#e2e8f0'))
    c.setLineWidth(0.7)
    c.roundRect(info_box_x, info_box_y, info_box_w, info_box_h, 2.5 * mm, fill=1, stroke=1)
    
    # Code Commercial
    c.setFont("Helvetica-Bold", 5.2)
    c.setFillColor(colors.HexColor('#64748b'))
    c.drawString(info_box_x + 2.5 * mm, info_box_y + 12.8 * mm, "CODE COMMERCIAL :")
    c.setFont("Helvetica-Bold", 7.5)
    c.setFillColor(colors.HexColor('#064e3b'))
    c.drawRightString(info_box_x + info_box_w - 2.5 * mm, info_box_y + 12.8 * mm, commercial['code'])
    
    # WhatsApp
    c.setFont("Helvetica-Bold", 5.2)
    c.setFillColor(colors.HexColor('#64748b'))
    c.drawString(info_box_x + 2.5 * mm, info_box_y + 7.5 * mm, "CONTACT WHATSAPP :")
    c.setFont("Helvetica-Bold", 6.8)
    c.setFillColor(colors.HexColor('#0f172a'))
    c.drawRightString(info_box_x + info_box_w - 2.5 * mm, info_box_y + 7.5 * mm, commercial['phone_display'])
    
    # Zone
    c.setFont("Helvetica-Bold", 5.2)
    c.setFillColor(colors.HexColor('#64748b'))
    c.drawString(info_box_x + 2.5 * mm, info_box_y + 2.2 * mm, "ZONE D'AFFECTATION :")
    c.setFont("Helvetica-Bold", 6.8)
    c.setFillColor(colors.HexColor('#0f172a'))
    c.drawRightString(info_box_x + info_box_w - 2.5 * mm, info_box_y + 2.2 * mm, commercial['zone'])
    
    # 7. Bandeau inférieur de sécurité
    footer_h = 6.8 * mm
    c.setFillColor(colors.HexColor('#064e3b'))
    c.roundRect(x + 0.6, y + 0.6, w - 1.2, footer_h, 3.5 * mm, fill=1, stroke=0)
    c.rect(x + 0.6, y + 3 * mm, w - 1.2, footer_h - 2.4 * mm, fill=1, stroke=0)
    
    c.setFont("Helvetica-Bold", 5.2)
    c.setFillColor(colors.white)
    c.drawCentredString(x + w / 2, y + 3.6 * mm, "CARTE OFFICIELLE FASOCARNET • 2026-2027")
    c.setFont("Helvetica", 4.2)
    c.setFillColor(colors.HexColor('#a7f3d0'))
    c.drawCentredString(x + w / 2, y + 1.6 * mm, "Validée par la Direction Commerciale & Support")
    
    c.restoreState()

def draw_badge_verso(c, x, y, w, h, commercial):
    c.saveState()
    # 1. Fond du badge
    c.setFillColor(colors.white)
    c.setStrokeColor(colors.HexColor('#064e3b'))
    c.setLineWidth(1.2)
    c.roundRect(x, y, w, h, 4.5 * mm, fill=1, stroke=1)
    
    # 2. Bandeau supérieur Verso
    header_h = 14 * mm
    c.setFillColor(colors.HexColor('#064e3b'))
    c.roundRect(x + 0.6, y + h - header_h - 0.6, w - 1.2, header_h, 4 * mm, fill=1, stroke=0)
    c.rect(x + 0.6, y + h - header_h - 0.6, w - 1.2, 5.5 * mm, fill=1, stroke=0)
    
    # Ligne accent or
    c.setFillColor(colors.HexColor('#d97706'))
    c.rect(x + 0.6, y + h - header_h - 1.4 * mm, w - 1.2, 1.2 * mm, fill=1, stroke=0)
    
    c.setFont("Helvetica-Bold", 8)
    c.setFillColor(colors.white)
    c.drawCentredString(x + w / 2, y + h - 6.5 * mm, "SCANNEZ POUR INSTALLER")
    
    c.setFont("Helvetica-Bold", 5.8)
    c.setFillColor(colors.HexColor('#fef08a'))
    c.drawCentredString(x + w / 2, y + h - 10.5 * mm, "L'APPLICATION FASOCARNET (PWA)")
    
    # 3. Zone QR Code
    qr_box_size = 29 * mm
    qr_box_x = x + (w - qr_box_size) / 2
    qr_box_y = y + 38 * mm
    
    c.setFillColor(colors.white)
    c.setStrokeColor(colors.HexColor('#cbd5e1'))
    c.setLineWidth(0.8)
    c.roundRect(qr_box_size_x := qr_box_x, qr_box_y, qr_box_size, qr_box_size, 2.5 * mm, fill=1, stroke=1)
    
    qr_size = 26 * mm
    qr_x = x + (w - qr_size) / 2
    qr_y = qr_box_y + 1.5 * mm
    qr_url = f"https://fasocarnet.onrender.com/?ref={commercial['code']}"
    draw_qr_code(c, qr_url, qr_x, qr_y, qr_size)
    
    # 4. Code Commercial sous le QR Code
    badge_ref_y = y + 31.5 * mm
    badge_ref_h = 5.2 * mm
    c.setFillColor(colors.HexColor('#ecfdf5'))
    c.setStrokeColor(colors.HexColor('#047857'))
    c.setLineWidth(0.8)
    c.roundRect(x + 3 * mm, badge_ref_y, w - 6 * mm, badge_ref_h, 1.8 * mm, fill=1, stroke=1)
    
    c.setFont("Helvetica-Bold", 5.8)
    c.setFillColor(colors.HexColor('#064e3b'))
    c.drawCentredString(x + w / 2, badge_ref_y + 1.4 * mm, f"CODE PARRAIN COMMERCIAL : {commercial['code']}")
    
    # 5. Les 3 étapes ultra simples
    step_box_y = y + 12.5 * mm
    step_box_h = 17.5 * mm
    step_box_w = w - 6 * mm
    step_box_x = x + 3 * mm
    
    c.setFillColor(colors.HexColor('#f8fafc'))
    c.setStrokeColor(colors.HexColor('#e2e8f0'))
    c.setLineWidth(0.7)
    c.roundRect(step_box_x, step_box_y, step_box_w, step_box_h, 2.5 * mm, fill=1, stroke=1)
    
    steps = [
        ("1", "Ouvrez l'appareil photo et pointez ce code."),
        ("2", "Touchez le lien web pour ouvrir la boutique."),
        ("3", "Appuyez sur « Installer » : l'icone apparait !")
    ]
    for i, (num, txt) in enumerate(steps):
        row_y = step_box_y + step_box_h - (i + 1) * 5.2 * mm + 1.1 * mm
        c.setFillColor(colors.HexColor('#047857'))
        c.circle(step_box_x + 3 * mm, row_y + 1.5 * mm, 1.8 * mm, fill=1, stroke=0)
        c.setFont("Helvetica-Bold", 4.6)
        c.setFillColor(colors.white)
        c.drawCentredString(step_box_x + 3 * mm, row_y + 0.6 * mm, num)
        
        c.setFont("Helvetica", 4.8)
        c.setFillColor(colors.HexColor('#334155'))
        c.drawString(step_box_x + 6.5 * mm, row_y + 0.4 * mm, txt)
        
    # 6. Puces de réassurance
    c.setFont("Helvetica-Bold", 4.8)
    c.setFillColor(colors.HexColor('#047857'))
    c.drawCentredString(x + w / 2, y + 8.8 * mm, "100% SANS INTERNET  •  MOINS DE 5 MO  •  SECURISE")
    
    # 7. Bandeau inférieur Support
    footer_h = 6.4 * mm
    c.setFillColor(colors.HexColor('#0f172a'))
    c.roundRect(x + 0.6, y + 0.6, w - 1.2, footer_h, 3.5 * mm, fill=1, stroke=0)
    c.rect(x + 0.6, y + 3 * mm, w - 1.2, footer_h - 2.4 * mm, fill=1, stroke=0)
    
    c.setFont("Helvetica-Bold", 4.8)
    c.setFillColor(colors.white)
    c.drawCentredString(x + w / 2, y + 2.4 * mm, "Assistance Commerciale & Technique : +226 72 99 03 10")
    
    c.restoreState()

def generate_badge_images(output_dir, commercial_list):
    """Génère les images PNG 300 DPI pour chaque commercial (recto et verso)"""
    os.makedirs(output_dir, exist_ok=True)
    badge_w = 54 * mm
    badge_h = 85.6 * mm
    scale = 300 / 72.0
    
    for comm in commercial_list:
        slug = comm['slug']
        
        # Recto
        pdf_r_path = os.path.join(output_dir, f"temp_r_{slug}.pdf")
        cr = canvas.Canvas(pdf_r_path, pagesize=(badge_w, badge_h))
        draw_badge_recto(cr, 0, 0, badge_w, badge_h, comm)
        cr.showPage()
        cr.save()
        
        pdf_r = pypdfium2.PdfDocument(pdf_r_path)
        img_r = pdf_r[0].render(scale=scale).to_pil()
        recto_png = os.path.join(output_dir, f"badge_recto_{slug}.png")
        img_r.save(recto_png, format="PNG")
        pdf_r.close()
        if os.path.exists(pdf_r_path):
            os.remove(pdf_r_path)
        comm['recto_img'] = recto_png
        
        # Verso
        pdf_v_path = os.path.join(output_dir, f"temp_v_{slug}.pdf")
        cv = canvas.Canvas(pdf_v_path, pagesize=(badge_w, badge_h))
        draw_badge_verso(cv, 0, 0, badge_w, badge_h, comm)
        cv.showPage()
        cv.save()
        
        pdf_v = pypdfium2.PdfDocument(pdf_v_path)
        img_v = pdf_v[0].render(scale=scale).to_pil()
        verso_png = os.path.join(output_dir, f"badge_verso_{slug}.png")
        img_v.save(verso_png, format="PNG")
        pdf_v.close()
        if os.path.exists(pdf_v_path):
            os.remove(pdf_v_path)
        comm['verso_img'] = verso_png
        has_photo = "[AVEC PHOTO]" if comm.get('photo_path') else "[SILHOUETTE]"
        print(f"[OK] Images generees pour {comm['name']} {has_photo}")

def set_cell_margins(cell, top=0, bottom=0, left=0, right=0):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for margin_name, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{margin_name}')
        node.set(qn('w:w'), str(int(val * 56.7)))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders_none(table):
    tblPr = table._tbl.tblPr
    tblBorders = OxmlElement('w:tblBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'none')
        tblBorders.append(border)
    tblPr.append(tblBorders)

def build_docx(commercial_list, output_path, desktop_path):
    doc = docx.Document()
    
    # Page A4 (210 mm x 297 mm)
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    
    margin_lr = 15 # mm
    margin_tb = 16 # mm
    section.left_margin = Mm(margin_lr)
    section.right_margin = Mm(margin_lr)
    section.top_margin = Mm(margin_tb)
    section.bottom_margin = Mm(margin_tb)
    
    badge_width = Mm(54.0)
    badge_height = Mm(85.6)
    col_width = Mm(90.0)
    
    # Découpage par lots de 4
    lots = []
    for i in range(0, len(commercial_list), 4):
        lots.append((commercial_list[i:i+4], (i // 4) + 1))
    
    for lot_idx, (groupe, batch_num) in enumerate(lots):
        if lot_idx > 0:
            doc.add_page_break()
            
        # =====================================================================
        # PAGE RECTO
        # =====================================================================
        p_head_r = doc.add_paragraph()
        p_head_r.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_head_r.paragraph_format.space_before = Pt(0)
        p_head_r.paragraph_format.space_after = Pt(6)
        r_head = p_head_r.add_run(f"FASOCARNET • BADGES ZONE ZINIARÉ • PLANCHE {batch_num} : RECTO (FACES AVANT)")
        r_head.bold = True
        r_head.font.size = Pt(9.5)
        r_head.font.color.rgb = RGBColor(6, 78, 59)
        
        table_r = doc.add_table(rows=2, cols=2)
        table_r.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders_none(table_r)
        
        # Remplir grille 2x2 (gérer le cas où il y a moins de 4 commerciaux dans le lot)
        grid_r = [
            [groupe[0] if len(groupe) > 0 else None, groupe[1] if len(groupe) > 1 else None],
            [groupe[2] if len(groupe) > 2 else None, groupe[3] if len(groupe) > 3 else None]
        ]
        
        for r_i in range(2):
            row = table_r.rows[r_i]
            row.height = Mm(115.0)
            for c_i in range(2):
                cell = row.cells[c_i]
                cell.width = col_width
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_margins(cell, top=2, bottom=4, left=0, right=0)
                
                comm = grid_r[r_i][c_i]
                if not comm:
                    continue
                p_text = cell.paragraphs[0]
                p_text.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_text.paragraph_format.space_before = Pt(0)
                p_text.paragraph_format.space_after = Pt(2)
                
                photo_status = "[Photo OK]" if comm.get('photo_path') else "[Silhouette]"
                r_title = p_text.add_run(f"● {comm['name']} | Code : {comm['code']} {photo_status}\n")
                r_title.font.size = Pt(7.5)
                r_title.font.bold = True
                r_title.font.color.rgb = RGBColor(100, 116, 139)
                
                p_img = cell.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(0)
                p_img.paragraph_format.space_after = Pt(0)
                r_img = p_img.add_run()
                r_img.add_picture(comm['recto_img'], width=badge_width, height=badge_height)
                
        p_foot_r = doc.add_paragraph()
        p_foot_r.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_foot_r.paragraph_format.space_before = Pt(6)
        p_foot_r.paragraph_format.space_after = Pt(0)
        r_foot = p_foot_r.add_run("Format officiel : 54 mm × 85,6 mm (Format Carte Bancaire / CR80) • Zone Ziniaré")
        r_foot.font.size = Pt(7.5)
        r_foot.font.italic = True
        r_foot.font.color.rgb = RGBColor(148, 163, 184)
        
        # =====================================================================
        # PAGE VERSO (DISPOSITION MIROIR POUR IMPRESSION RECTO-VERSO)
        # =====================================================================
        doc.add_page_break()
        
        p_head_v = doc.add_paragraph()
        p_head_v.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_head_v.paragraph_format.space_before = Pt(0)
        p_head_v.paragraph_format.space_after = Pt(6)
        r_head_v = p_head_v.add_run(f"FASOCARNET • BADGES ZONE ZINIARÉ • PLANCHE {batch_num} : VERSO (QR CODES)")
        r_head_v.bold = True
        r_head_v.font.size = Pt(9.5)
        r_head_v.font.color.rgb = RGBColor(6, 78, 59)
        
        table_v = doc.add_table(rows=2, cols=2)
        table_v.alignment = WD_TABLE_ALIGNMENT.CENTER
        set_table_borders_none(table_v)
        
        # Disposition miroir horizontal
        grid_v = [
            [groupe[1] if len(groupe) > 1 else None, groupe[0] if len(groupe) > 0 else None],
            [groupe[3] if len(groupe) > 3 else None, groupe[2] if len(groupe) > 2 else None]
        ]
        
        for r_i in range(2):
            row = table_v.rows[r_i]
            row.height = Mm(115.0)
            for c_i in range(2):
                cell = row.cells[c_i]
                cell.width = col_width
                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                set_cell_margins(cell, top=2, bottom=4, left=0, right=0)
                
                comm = grid_v[r_i][c_i]
                if not comm:
                    continue
                p_text = cell.paragraphs[0]
                p_text.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_text.paragraph_format.space_before = Pt(0)
                p_text.paragraph_format.space_after = Pt(2)
                
                r_title = p_text.add_run(f"● Verso QR • {comm['name']} ({comm['code']})\n")
                r_title.font.size = Pt(7.5)
                r_title.font.bold = True
                r_title.font.color.rgb = RGBColor(100, 116, 139)
                
                p_img = cell.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(0)
                p_img.paragraph_format.space_after = Pt(0)
                r_img = p_img.add_run()
                r_img.add_picture(comm['verso_img'], width=badge_width, height=badge_height)
                
        p_foot_v = doc.add_paragraph()
        p_foot_v.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p_foot_v.paragraph_format.space_before = Pt(6)
        p_foot_v.paragraph_format.space_after = Pt(0)
        r_foot_v = p_foot_v.add_run("Impression recto-verso : sélectionner « Retourner sur les bords longs » (Flip on long edge)")
        r_foot_v.font.size = Pt(7.5)
        r_foot_v.font.italic = True
        r_foot_v.font.color.rgb = RGBColor(148, 163, 184)
        
    doc.save(output_path)
    print(f"[OK] Fichier docx enregistre : {output_path}")
    
    # Copie sur le Bureau
    try:
        doc.save(desktop_path)
        print(f"[OK] Fichier copie sur le Bureau : {desktop_path}")
    except Exception as e:
        print(f"[AVERTISSEMENT] Erreur copie Bureau : {e}")

if __name__ == '__main__':
    images_dir = os.path.join(os.getcwd(), "docs", "badges_ziniare")
    
    # Inclure Hyppolite s'il s'agit d'un 9ème commercial pour Ziniaré
    all_commercials = list(COMMERCIAUX_ZINIARE)
    # Ajouter Hyppolite en 9ème
    all_commercials.append(COMMERCIAL_HYPPOLITE)
    
    generate_badge_images(images_dir, all_commercials)
    
    out_docx = os.path.join(os.getcwd(), "docs", "badges_commerciaux_ziniare.docx")
    public_docx = os.path.join(os.getcwd(), "public", "badges_commerciaux_ziniare.docx")
    desktop_docx = r"C:\Users\Maxime OUATTARA\Desktop\badges_commerciaux_ziniare.docx"
    
    build_docx(all_commercials, out_docx, desktop_docx)
    
    # Copie aussi dans public
    shutil.copyfile(out_docx, public_docx)
    print(f"[OK] Fichier copie dans public : {public_docx}")
