import os
import sys
import docx
from docx.shared import Mm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_margins(cell, top=0, bottom=0, left=0, right=0):
    """Définit les marges internes (padding) d'une cellule en dxa (1 mm = 56.7 dxa)"""
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for margin_name, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{margin_name}')
        node.set(qn('w:w'), str(int(val * 56.7)))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_table_borders_none(table):
    """Enlève toutes les bordures d'un tableau Word pour une mise en page épurée"""
    tblPr = table._tbl.tblPr
    tblBorders = OxmlElement('w:tblBorders')
    for border_name in ['top', 'left', 'bottom', 'right', 'insideH', 'insideV']:
        border = OxmlElement(f'w:{border_name}')
        border.set(qn('w:val'), 'none')
        tblBorders.append(border)
    tblPr.append(tblBorders)

def create_badges_docx():
    doc = docx.Document()
    
    # Configuration de la page A4 (210 mm x 297 mm)
    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    
    # Marges symétriques pour un alignement recto-verso parfait
    margin_lr = 15 # mm
    margin_tb = 16 # mm
    section.left_margin = Mm(margin_lr)
    section.right_margin = Mm(margin_lr)
    section.top_margin = Mm(margin_tb)
    section.bottom_margin = Mm(margin_tb)
    
    # Dossiers sources des images
    docs_dir = os.path.join(os.getcwd(), "docs")
    public_dir = os.path.join(os.getcwd(), "public")
    artifact_dir = r"C:\Users\Maxime OUATTARA\.gemini\antigravity\brain\23e09812-2870-45e7-9025-1667cd904de7"
    
    # 4 Commerciaux
    commerciaux = [
        {
            'name': 'COMPAORÉ Adama',
            'code': 'COMPAORE226',
            'recto': os.path.join(docs_dir, "badge_recto_COMPAORE_Adama.png"),
            'verso': os.path.join(docs_dir, "badge_verso_COMPAORE_Adama.png")
        },
        {
            'name': 'SARÉ Falilatou',
            'code': 'SARE226',
            'recto': os.path.join(docs_dir, "badge_recto_SARE_Falilatou.png"),
            'verso': os.path.join(docs_dir, "badge_verso_SARE_Falilatou.png")
        },
        {
            'name': 'SARÉ Salamatou',
            'code': 'SARE7',
            'recto': os.path.join(docs_dir, "badge_recto_SARE_Salamatou.png"),
            'verso': os.path.join(docs_dir, "badge_verso_SARE_Salamatou.png")
        },
        {
            'name': 'SOGLI Rebécca',
            'code': 'SOGLI226',
            'recto': os.path.join(docs_dir, "badge_recto_SOGLI_Rebecca.png"),
            'verso': os.path.join(docs_dir, "badge_verso_SOGLI_Rebecca.png")
        }
    ]
    
    # Taille exacte du badge CR80 standard
    badge_width = Mm(54.0)
    badge_height = Mm(85.6)
    
    # Largeur de chaque colonne (180 mm / 2 = 90 mm)
    col_width = Mm(90.0)
    
    # =========================================================================
    # PAGE 1 : RECTO (FACES AVANT)
    # =========================================================================
    
    # En-tête discret
    p_header1 = doc.add_paragraph()
    p_header1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_header1.paragraph_format.space_before = Pt(0)
    p_header1.paragraph_format.space_after = Pt(8)
    run1 = p_header1.add_run("FASOCARNET • BADGES COMMERCIAUX TERRAIN • PAGE 1 : RECTO")
    run1.bold = True
    run1.font.size = Pt(9.5)
    run1.font.color.rgb = RGBColor(6, 78, 59) # Vert émeraude
    
    # Tableau 2 colonnes x 2 lignes pour le Recto
    table_recto = doc.add_table(rows=2, cols=2)
    table_recto.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders_none(table_recto)
    
    # Grille Recto :
    # (0, 0) : COMPAORE Adama | (0, 1) : SARE Falilatou
    # (1, 0) : SARE Salamatou | (1, 1) : SOGLI Rebecca
    recto_grid = [
        [commerciaux[0], commerciaux[1]],
        [commerciaux[2], commerciaux[3]]
    ]
    
    for r_idx in range(2):
        row = table_recto.rows[r_idx]
        row.height = Mm(115.0) # Hauteur suffisante pour badge 85.6 mm + repère
        for c_idx in range(2):
            cell = row.cells[c_idx]
            cell.width = col_width
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=2, bottom=4, left=0, right=0)
            
            comm = recto_grid[r_idx][c_idx]
            
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(3)
            
            # Petit repère textuel discret
            r_title = p.add_run(f"● {comm['name']} ({comm['code']})\n")
            r_title.font.size = Pt(7.5)
            r_title.font.bold = True
            r_title.font.color.rgb = RGBColor(100, 116, 139)
            
            # Insertion image Recto à taille exacte
            p_img = cell.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(0)
            p_img.paragraph_format.space_after = Pt(0)
            r_img = p_img.add_run()
            r_img.add_picture(comm['recto'], width=badge_width, height=badge_height)
            
    # Pied de page 1 informatif
    p_foot1 = doc.add_paragraph()
    p_foot1.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_foot1.paragraph_format.space_before = Pt(8)
    p_foot1.paragraph_format.space_after = Pt(0)
    r_foot1 = p_foot1.add_run("Format officiel carte : 54 mm × 85,6 mm • Découpage le long du contour arrondi")
    r_foot1.font.size = Pt(7.5)
    r_foot1.font.italic = True
    r_foot1.font.color.rgb = RGBColor(148, 163, 184)
    
    # =========================================================================
    # PAGE 2 : VERSO (FACES ARRIÈRE / QR CODES)
    # =========================================================================
    doc.add_page_break()
    
    # En-tête identique pour garantir une hauteur de départ rigoureusement égale
    p_header2 = doc.add_paragraph()
    p_header2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_header2.paragraph_format.space_before = Pt(0)
    p_header2.paragraph_format.space_after = Pt(8)
    run2 = p_header2.add_run("FASOCARNET • BADGES COMMERCIAUX TERRAIN • PAGE 2 : VERSO (QR)")
    run2.bold = True
    run2.font.size = Pt(9.5)
    run2.font.color.rgb = RGBColor(6, 78, 59)
    
    # Tableau 2 colonnes x 2 lignes pour le Verso
    table_verso = doc.add_table(rows=2, cols=2)
    table_verso.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_borders_none(table_verso)
    
    # Disposition miroir horizontal pour que l'impression recto-verso (retournement bord long)
    # superpose exactement chaque verso sur son recto respectif :
    # Recto (0, 0) [Adama]  <--> Verso (0, 1) [Adama]
    # Recto (0, 1) [Falilatou] <--> Verso (0, 0) [Falilatou]
    # Recto (1, 0) [Salamatou] <--> Verso (1, 1) [Salamatou]
    # Recto (1, 1) [Rebecca]  <--> Verso (1, 0) [Rebecca]
    verso_grid = [
        [commerciaux[1], commerciaux[0]],
        [commerciaux[3], commerciaux[2]]
    ]
    
    for r_idx in range(2):
        row = table_verso.rows[r_idx]
        row.height = Mm(115.0)
        for c_idx in range(2):
            cell = row.cells[c_idx]
            cell.width = col_width
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=2, bottom=4, left=0, right=0)
            
            comm = verso_grid[r_idx][c_idx]
            
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(3)
            
            # Petit repère textuel discret
            r_title = p.add_run(f"● Verso QR • {comm['name']} ({comm['code']})\n")
            r_title.font.size = Pt(7.5)
            r_title.font.bold = True
            r_title.font.color.rgb = RGBColor(100, 116, 139)
            
            # Insertion image Verso à taille exacte
            p_img = cell.add_paragraph()
            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p_img.paragraph_format.space_before = Pt(0)
            p_img.paragraph_format.space_after = Pt(0)
            r_img = p_img.add_run()
            r_img.add_picture(comm['verso'], width=badge_width, height=badge_height)
            
    # Pied de page 2
    p_foot2 = doc.add_paragraph()
    p_foot2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_foot2.paragraph_format.space_before = Pt(8)
    p_foot2.paragraph_format.space_after = Pt(0)
    r_foot2 = p_foot2.add_run("Impression recto-verso : sélectionner « Retourner sur les bords longs » (Flip on long edge)")
    r_foot2.font.size = Pt(7.5)
    r_foot2.font.italic = True
    r_foot2.font.color.rgb = RGBColor(148, 163, 184)
    
    # Enregistrement du fichier Word
    docx_path = os.path.join(docs_dir, "badges_commerciaux_fasocarnet.docx")
    public_docx = os.path.join(public_dir, "badges_commerciaux_fasocarnet.docx")
    doc.save(docx_path)
    doc.save(public_docx)
    
    if os.path.exists(artifact_dir):
        doc.save(os.path.join(artifact_dir, "badges_commerciaux_fasocarnet.docx"))
        
    print(f"[OK] Fichier Word genere avec succes : {docx_path}")

if __name__ == '__main__':
    create_badges_docx()
